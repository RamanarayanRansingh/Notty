import streamlit as st
from langchain_community.document_loaders import WebBaseLoader
from chains.code_analyzer import CodeAnalyzer
from chains.docs_analyzer import DocsAnalyzer
from utils.code_processor import CodeProcessor
from utils.docs_processor import DocsProcessor
from config import Config
# Required libraries for export formats
from io import BytesIO
from fpdf import FPDF  # For PDF export
from docx import Document  # For DOCX export

# Initialize components
code_analyzer = CodeAnalyzer()
docs_analyzer = DocsAnalyzer()

# Page config
st.set_page_config(
    layout="wide",
    page_title="Tech Documentation & Code Analyzer",
    page_icon="🔍"
)

# Custom CSS for styling
st.markdown("""
<style>
.code-explanation {
    background-color: #1E1E1E;
    color: white;
    padding: 1rem;
    border-radius: 5px;
    margin: 1rem 0;
    overflow-x: auto;
}
.concept-card {
    background-color: #2C3E50;
    color: white;
    padding: 1rem;
    border-radius: 5px;
    margin: 0.5rem 0;
}
.code-block {
    background-color: #2d2d2d;
    color: #f8f8f2;
    font-family: monospace;
    padding: 10px;
    border-radius: 5px;
    margin: 10px 0;
    overflow-x: auto;
}
</style>
""", unsafe_allow_html=True)

# State management for analysis results
if "doc_analysis" not in st.session_state:
    st.session_state.doc_analysis = None
if "doc_content" not in st.session_state:
    st.session_state.doc_content = None

# Sidebar navigation
st.sidebar.title("Navigation")
page = st.sidebar.radio(
    "Choose functionality",
    ["Documentation Analyzer", "Code Analyzer"]
)

# Documentation Analyzer Page
if page == "Documentation Analyzer":
    st.title("📚 Documentation Analyzer")
    
    # Input method for documentation
    input_method = st.radio("Choose input method", ["URL", "Text Input"])
    
    if input_method == "URL":
        url = st.text_input("Enter documentation URL")
        if url and st.button("Analyze Documentation"):
            with st.spinner("Analyzing documentation..."):
                try:
                    loader = WebBaseLoader([url])
                    content = loader.load().pop().page_content
                    cleaned_content = DocsProcessor.clean_content(content)
                    st.session_state.doc_content = cleaned_content
                    
                    # Analyze documentation
                    analysis = docs_analyzer.analyze_documentation(cleaned_content)
                    st.session_state.doc_analysis = analysis
                    
                    st.success("Documentation analyzed successfully!")
                except Exception as e:
                    st.error(f"Error analyzing documentation: {e}")
    
    elif input_method == "Text Input":
        text_input = st.text_area("Paste your documentation text here")
        if text_input and st.button("Analyze Documentation"):
            with st.spinner("Analyzing documentation..."):
                try:
                    cleaned_content = DocsProcessor.clean_content(text_input)
                    st.session_state.doc_content = cleaned_content
                    
                    analysis = docs_analyzer.analyze_documentation(cleaned_content)
                    st.session_state.doc_analysis = analysis
                    
                    st.success("Documentation analyzed successfully!")
                except Exception as e:
                    st.error(f"Error analyzing documentation: {e}")
    
# Display results
if st.session_state.doc_analysis:
    st.header("📋 Analysis Results")
    analysis = st.session_state.doc_analysis

    if isinstance(analysis, dict):  # Ensure it's a dictionary
        # Display main topic and prerequisites
        col1, col2 = st.columns(2)
        with col1:
            st.subheader("Main Topic")
            st.write(analysis.get("main_topic", "N/A"))
        with col2:
            st.subheader("Prerequisites")
            for prereq in analysis.get("prerequisites", []):
                st.write(f"• {prereq}")
        
        # Display key concepts
        st.subheader("Key Concepts")
        for concept in analysis.get("key_concepts", []):
            st.markdown(f'<div class="concept-card">{concept}</div>', unsafe_allow_html=True)
        
        # Include code snippets
        st.subheader("Code Snippets")
        for code in analysis.get("code_snippets", []):
            st.markdown(f'<div class="code-block">{code}</div>', unsafe_allow_html=True)
    else:
        # If analysis is a string, display it directly
        st.write("Analysis Results:")
        st.markdown(analysis)

        
        # Generate and display notes
        st.subheader("📒 Notes")
        notes = docs_analyzer.generate_notes(st.session_state.doc_content)
        st.markdown(notes, unsafe_allow_html=True)

        # Export options
st.subheader("Export Notes")
export_format = st.selectbox(
    "Select format",
    ["HTML", "Markdown", "PDF", "DOCX"],  # Include all options
    key="export_format"
)

if st.button("Export Notes", key="export_notes"):
    try:
        # Format notes based on the selected format
        if export_format == "HTML":
            formatted_notes = DocsProcessor.format_for_export(notes, "HTML")
            mime_type = "text/html"
            file_name = "notes.html"
        elif export_format == "Markdown":
            formatted_notes = DocsProcessor.format_for_export(notes, "Markdown")
            mime_type = "text/markdown"
            file_name = "notes.md"
        elif export_format == "PDF":
            # Create a PDF using FPDF
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            for line in notes.split("\n"):
                pdf.cell(200, 10, txt=line, ln=True, align='L')
            # Output PDF as bytes
            formatted_notes = pdf.output(dest="S").encode("latin1")
            mime_type = "application/pdf"
            file_name = "notes.pdf"
        elif export_format == "DOCX":
            # Create a DOCX using python-docx
            doc = Document()
            for line in notes.split("\n"):
                doc.add_paragraph(line)
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            formatted_notes = buffer.read()
            mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            file_name = "notes.docx"
        
        # Provide the download button
        st.download_button(
            label=f"Download Notes as {export_format.upper()}",
            data=formatted_notes,
            file_name=file_name,
            mime=mime_type
        )
    except Exception as e:
        st.error(f"Error exporting notes: {e}")

# Code Analyzer Page
elif page == "Code Analyzer":
    st.title("💻 Code Analyzer")
    code_input = st.text_area("Paste your code here")
    language = st.selectbox("Select programming language", Config.SUPPORTED_LANGUAGES)
    analysis_type = st.radio("Choose analysis type", ["Full Analysis", "Explain Specific Lines"])
    
    if code_input and st.button("Analyze Code"):
        with st.spinner("Analyzing code..."):
            try:
                if analysis_type == "Full Analysis":
                    analysis = code_analyzer.analyze_code_block(code_input, language)
                    st.subheader("🔍 Code Analysis")
                    st.markdown(f'<div class="code-explanation">{analysis}</div>', unsafe_allow_html=True)
                elif analysis_type == "Explain Specific Lines":
                    line_numbers = st.text_input("Enter line numbers (e.g., 1-3, 5)")
                    if line_numbers:
                        parsed_lines = CodeProcessor.parse_line_numbers(line_numbers)
                        analysis = code_analyzer.explain_specific_lines(code_input, parsed_lines)
                        st.subheader("📝 Explanation for Specific Lines")
                        st.markdown(f'<div class="code-explanation">{analysis}</div>', unsafe_allow_html=True)
            except Exception as e:
                st.error(f"Error analyzing code: {e}")
